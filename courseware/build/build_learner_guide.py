#!/usr/bin/env python3
"""Generate the CompTIA Security+ (SY0-701) Learner Guide as BOTH a Markdown
mirror (LG-*.md at repo root) and a DOCX (courseware/LG-*.docx) from one source,
so they never diverge.

House format: cover page, Document Version Control Record, auto TOC, Arial 11pt
body, one section per lab (Objective · Goal · What you'll build · Step-by-step
with commands · Test it), plus setup, exam-prep and glossary. All content is
driven by course_data + the domain data files, keeping the LG 100% aligned with
the slide deck, Lesson Plan and labs.
"""
import os, sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3; from data_domain4 import DOMAIN4
from data_domain5 import DOMAIN5
ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4+DOMAIN5
import prodoc
COURSEWARE=os.path.dirname(HERE); PROJ=os.path.dirname(COURSEWARE)
ASSETS=os.path.join(COURSEWARE,"assets")

# ---------------- block DSL (single content stream → MD + DOCX) ----------------
B=[]
def h1(t): B.append(("h1",t))
def h2(t): B.append(("h2",t))
def h3(t): B.append(("h3",t))
def p(t):  B.append(("p",t))
def bullets(xs): B.append(("bullets",xs))
def steps(xs): B.append(("steps",xs))
def code(t): B.append(("code",t))
def note(t): B.append(("note",t))
def rule(): B.append(("rule",))

# ---------------- content ----------------
h1("Introduction")
p(f"This Learner Guide accompanies the WSQ course {C.TITLE} ({C.COURSE_CODE}), conducted by {C.ORG}. "
  "It provides step-by-step instructions for all 36 hands-on labs, organised by the five official "
  "CompTIA Security+ SY0-701 exam domains. Every lab maps to a published exam objective and runs "
  "in the free Killercoda Ubuntu Playground — a disposable Linux VM in your browser.")
p("Use this guide alongside the course slides and the full lab files in the labs/ folder of the course "
  f"repository ({C.REPO_URL}). Each lab notes the free tools it uses and how to verify success; reset "
  "the Killercoda playground between labs that change firewall, PAM or routing state.")

h1("Course Learning Outcomes")
bullets(C.LEARNING_OUTCOMES)

h1("Before You Start — Environment Setup")
h3("What you need")
bullets([
 "A modern web browser — no local installation is required.",
 f"The free Killercoda Ubuntu Playground: {C.KILLERCODA_URL}",
 "A root shell with internet access (provided by Killercoda) to install tools via apt or a single binary download.",
 "The course lab files — clone or download the repository at "+C.REPO_URL+".",
])
h3("Launch the Killercoda Ubuntu Playground")
p("Open the playground in your browser and wait for the terminal to become ready. You are dropped into a "
  "root shell on a throwaway Ubuntu VM, so you can install packages and change system state freely. When a "
  "lab is finished — or if a lab leaves the VM in a bad state — click Restart/Reset to get a clean VM.")
code("whoami            # root\nlsb_release -a    # confirm the Ubuntu version\napt update        # refresh the package lists before installing lab tools")
h3("Conventions used in every lab")
bullets([
 "Each lab lists its Required software at the top; install it with the apt or download command shown.",
 "Commands are run as root in the Killercoda terminal unless a lab says otherwise.",
 "Reset the playground between labs that change firewall, PAM, routing or disk state.",
 "A few labs also use free online tools (e.g. VirusTotal, SSL Labs, crt.sh) — links are in each lab file.",
])
p("Legal & ethical use: labs involving scanning, exploitation, password cracking, phishing simulation or "
  "denial-of-service (Labs 7, 9, 11, 32, 33, 34, 36) must only be run against systems you own (the "
  "Killercoda VM, your own containers) or targets you are explicitly authorised to test. Unauthorised use "
  "against third-party systems is illegal, including under the Singapore Computer Misuse Act (Cap. 50A).")

# ---------------- per-domain, per-lab ----------------
for t in C.TOPICS:
    h1(f"Domain {t['code']} — {t['title']}  ({t['weighting']})")
    p(t["subtitle"])
    h3("Key concepts")
    bullets(t["concepts"])
    for a in [x for x in ACT if x["topic"]==t["num"]]:
        h2(f"Lab {a['num']} — {a['title']}")
        p(f"Exam objective: {a['objective']}.")
        p(f"Goal: {a['desc']}")
        h3("What you'll build")
        p(a["build"]+f"   (Tools: {a['services']}.)")
        h3("Step-by-step")
        steps([(instr,cmd) for (instr,cmd) in a["steps"]])
        h3("Test it")
        p(a["test"])
        note(f"Full step-by-step commands and references are in labs/lab-{a['num']:02d}-*.md. "
             "Reset the Killercoda playground before the next stateful lab.")
        rule()

h1("Exam Preparation")
bullets([
 "First pass: do every lab in Killercoda, reading the References in each lab file.",
 "Second pass: redo the labs from memory until the commands and concepts are automatic.",
 "Review the 'Test it' check and the 'What you learned' notes for any topic you find hard.",
 "Memorise the control categories/types, the CIA triad and AAA, and the CVSS base metrics.",
 f"Take the Tertiary Infotech Security+ practice exam: {C.PRACTICE_EXAM_URL}",
 "The SY0-701 exam is up to 90 questions in 90 minutes; the passing score is 750/900. Book it through Pearson VUE.",
])

h1("Glossary")
gl=[
 ("Security control","A safeguard classed by category (technical, managerial, operational, physical) and type (preventive, deterrent, detective, corrective, compensating, directive)."),
 ("CIA triad","Confidentiality, Integrity, Availability — the three goals of information security."),
 ("AAA","Authentication (who you are), Authorization (what you may do), Accounting (what you did)."),
 ("Hashing","A one-way function (e.g. SHA-256) that proves integrity; salting defeats rainbow tables."),
 ("Symmetric encryption","One shared secret key encrypts and decrypts (e.g. AES) — fast for bulk data."),
 ("Asymmetric encryption","A public/private key pair (e.g. RSA) that solves key exchange and enables signatures."),
 ("PKI","Public Key Infrastructure — certificate authorities bind public keys to verified identities."),
 ("SPF / DKIM / DMARC","Email authentication standards that expose spoofed and phishing mail."),
 ("SQL injection (SQLi)","Injecting SQL through unvalidated input; mitigated by parameterised queries."),
 ("Cross-site scripting (XSS)","Injecting script into a page; mitigated by output encoding and CSP."),
 ("WAF","Web Application Firewall — filters malicious HTTP requests (e.g. ModSecurity + OWASP CRS)."),
 ("SIEM","Security Information and Event Management — centralises and correlates logs and raises alerts."),
 ("IDS","Intrusion Detection System — inspects network traffic against signatures (e.g. Suricata)."),
 ("EDR","Endpoint Detection and Response — interrogates endpoint state and activity (e.g. osquery)."),
 ("FIM","File Integrity Monitoring — detects unauthorised file changes (e.g. AIDE)."),
 ("CVSS","Common Vulnerability Scoring System — a 0–10 severity score for a vulnerability."),
 ("Penetration test","An authorised, scoped simulated attack that validates real exploitability."),
]
B.append(("dl",gl))

# ---------------- render Markdown ----------------
def _anchor(txt):
    return "".join(ch.lower() if ch.isalnum() else ("-" if ch in " -" else "") for ch in txt)

def render_md():
    out=[f"# {C.TITLE} — Learner Guide",""]
    out.append(f"**WSQ Course Code:** {C.COURSE_CODE}  |  **Conducted by:** {C.ORG} ({C.UEN.replace('UEN: ','UEN ')})  |  **Version {C.VERSION} · {C.VERSION_DATE}**")
    out.append("")
    out.append("## Contents"); out.append("")
    for kind,*rest in B:
        if kind=="h1": out.append(f"- [{rest[0]}](#{_anchor(rest[0])})")
        elif kind=="h2": out.append(f"  - [{rest[0]}](#{_anchor(rest[0])})")
    out.append("")
    for kind,*rest in B:
        if kind=="h1": out+=["",f"## {rest[0]}",""]
        elif kind=="h2": out+=["",f"### {rest[0]}",""]
        elif kind=="h3": out+=[f"**{rest[0]}**",""]
        elif kind=="p": out+=[rest[0],""]
        elif kind=="bullets": out+=[f"- {x}" for x in rest[0]]+[""]
        elif kind=="steps":
            for i,(instr,cmd) in enumerate(rest[0],1):
                out.append(f"{i}. {instr}")
                if cmd: out+=["",f"   ```bash",f"   {cmd}","   ```",""]
            out.append("")
        elif kind=="code": out+=["```bash",rest[0],"```",""]
        elif kind=="note": out+=[f"> **Note:** {rest[0]}",""]
        elif kind=="rule": out+=["---",""]
        elif kind=="dl":
            for term,defn in rest[0]: out.append(f"- **{term}** — {defn}")
            out.append("")
    return "\n".join(out)

MD_OUT=os.path.join(PROJ,f"LG-{C.SHORT_TITLE}.md")
with open(MD_OUT,"w") as f: f.write(render_md())
print("Saved",MD_OUT)

# ---------------- render DOCX ----------------
BRAND=RGBColor(0x1F,0x6F,0xEB); DARK=RGBColor(0x11,0x18,0x27); GREY=RGBColor(0x55,0x5B,0x66)
INKCODE=RGBColor(0x0B,0x30,0x60)
doc=Document()
normal=doc.styles["Normal"]; normal.font.name="Arial"; normal.font.size=Pt(11)
prodoc.style_headings(doc)
prodoc.add_cover_page(doc,"LEARNER GUIDE",C.TITLE,C.VERSION.lstrip("v"),
                      org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),
                      course_logo=None, course_code=C.COURSE_CODE)
prodoc.add_version_control(doc,[(C.VERSION.lstrip("v"),C.VERSION_DATE,"Initial release — Security+ SY0-701 Learner Guide covering all 36 labs.",C.TRAINER)])
prodoc.add_toc(doc)

def code_para(text):
    for line in text.split("\n"):
        para=doc.add_paragraph(); prodoc._shade_para(para) if hasattr(prodoc,"_shade_para") else None
        r=para.add_run(line); r.font.name="Consolas"; r.font.size=Pt(9.5); r.font.color.rgb=INKCODE

for kind,*rest in B:
    if kind=="h1": doc.add_heading(rest[0],level=1)
    elif kind=="h2": doc.add_heading(rest[0],level=2)
    elif kind=="h3":
        para=doc.add_paragraph(); r=para.add_run(rest[0]); r.bold=True; r.font.size=Pt(11); r.font.color.rgb=BRAND
    elif kind=="p": doc.add_paragraph(rest[0])
    elif kind=="bullets":
        for x in rest[0]: doc.add_paragraph(x,style="List Bullet")
    elif kind=="steps":
        for i,(instr,cmd) in enumerate(rest[0],1):
            para=doc.add_paragraph(style="List Number"); para.add_run(instr)
            if cmd: code_para(cmd)
    elif kind=="code": code_para(rest[0])
    elif kind=="note":
        para=doc.add_paragraph(); r=para.add_run("Note: "); r.bold=True; r.font.color.rgb=BRAND
        para.add_run(rest[0]).font.size=Pt(10)
    elif kind=="rule": doc.add_paragraph("")
    elif kind=="dl":
        for term,defn in rest[0]:
            para=doc.add_paragraph(style="List Bullet")
            r=para.add_run(term+" — "); r.bold=True; para.add_run(defn)

prodoc.add_page_numbers(doc)
prodoc.enable_update_fields(doc)
DOCX_OUT=os.path.join(COURSEWARE,f"LG-{C.SHORT_TITLE}.docx")
doc.save(DOCX_OUT)
print("Saved",DOCX_OUT)
