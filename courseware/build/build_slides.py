#!/usr/bin/env python3
"""Generate the CompTIA Security+ (SY0-701) course slide deck (all-white
Tertiary house style).

Visual component library (cover, section, content, two_col, cards3,
tile_grid, flow_h, trainer_slide, big_statement, activity_overview,
step_slide, test_slide, brk) is the shared WSQ set. Content is driven entirely
by course_data.py + data_domainN.py so the deck stays 100% aligned with the LP,
LG and labs.
"""
import os, sys, math
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import course_data as C
from data_domain1 import DOMAIN1
from data_domain2 import DOMAIN2
from data_domain3 import DOMAIN3
from data_domain4 import DOMAIN4
from data_domain5 import DOMAIN5
ACTIVITIES = DOMAIN1 + DOMAIN2 + DOMAIN3 + DOMAIN4 + DOMAIN5

REPO = os.path.dirname(HERE)                      # courseware/
ASSETS = os.path.join(REPO, "assets")

# ---------------- palette ----------------
BLUE=RGBColor(0x1F,0x6F,0xEB); TEAL=RGBColor(0x10,0xB9,0x81); AMBER=RGBColor(0xF5,0x9E,0x0B)
INK=RGBColor(0x16,0x1B,0x26); GREY=RGBColor(0x5B,0x63,0x72); LIGHT=RGBColor(0xF5,0xF8,0xFC)
WHITE=RGBColor(0xFF,0xFF,0xFF); LINE=RGBColor(0xE2,0xE8,0xF0); VIOLET=RGBColor(0x7C,0x3A,0xED)
PALETTE=[BLUE,TEAL,VIOLET,AMBER]

prs=Presentation(); prs.slide_width=Inches(13.333); prs.slide_height=Inches(7.5)
SW,SH=prs.slide_width,prs.slide_height
BLANK=prs.slide_layouts[6]

def slide(): return prs.slides.add_slide(BLANK)
def rect(s,x,y,w,h,color,line=None):
    sp=s.shapes.add_shape(1,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=color
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb=line; sp.line.width=Pt(1)
    sp.shadow.inherit=False; return sp
def oval(s,x,y,w,h,color):
    sp=s.shapes.add_shape(9,x,y,w,h); sp.fill.solid(); sp.fill.fore_color.rgb=color
    sp.line.fill.background(); sp.shadow.inherit=False; return sp
def txt(s,x,y,w,h,runs,align=PP_ALIGN.LEFT,anchor=MSO_ANCHOR.TOP,space=4):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True; tf.vertical_anchor=anchor
    for i,line in enumerate(runs):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.alignment=align; p.space_after=Pt(space)
        for t,sz,col,bold in line:
            r=p.add_run(); r.text=t; r.font.size=Pt(sz); r.font.bold=bold
            r.font.color.rgb=col; r.font.name="Arial"
    return tb
def bullets(s,x,y,w,h,items,size=18,color=INK,gap=10,mcolor=BLUE):
    tb=s.shapes.add_textbox(x,y,w,h); tf=tb.text_frame; tf.word_wrap=True
    for i,it in enumerate(items):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.space_after=Pt(gap)
        lvl=it[1] if isinstance(it,tuple) else 0
        text=it[0] if isinstance(it,tuple) else it
        r=p.add_run(); r.text=("•  " if lvl==0 else "–  ")+text
        r.font.size=Pt(size if lvl==0 else size-2); r.font.color.rgb=color if lvl==0 else GREY
        r.font.name="Arial"; r.font.bold=(lvl==0 and isinstance(it,tuple) and len(it)>2 and it[2])
    return tb

PAGE={"n":0}
def footer(s):
    PAGE["n"]+=1
    txt(s,Inches(0.4),Inches(7.05),Inches(7.5),Inches(0.35),
        [[(f"{C.SHORT_TITLE}  ·  {C.COURSE_CODE}",9,GREY,False)]])
    txt(s,Inches(5.0),Inches(7.05),Inches(3.3),Inches(0.35),
        [[("© 2026 Tertiary Infotech Academy Pte Ltd",9,GREY,False)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(12.4),Inches(7.05),Inches(0.6),Inches(0.35),
        [[(str(PAGE["n"]),9,GREY,False)]],align=PP_ALIGN.RIGHT)
def head(s,title,kicker=None,kcolor=BLUE):
    rect(s,0,0,SW,SH,WHITE); rect(s,0,0,Inches(0.28),Inches(1.55),kcolor)
    if kicker: txt(s,Inches(0.85),Inches(0.5),Inches(11.6),Inches(0.4),[[(kicker,14,kcolor,True)]])
    txt(s,Inches(0.85),Inches(0.9),Inches(11.9),Inches(0.9),[[(title,29,INK,True)]])
    rect(s,Inches(0.85),Inches(1.7),Inches(11.63),Inches(0.02),LINE)
    return s
def _logo(name):
    p=os.path.join(ASSETS,name)
    return p if os.path.exists(p) else None

# ---------------- slide templates ----------------
def cover():
    s=slide(); rect(s,0,0,SW,SH,WHITE)
    rect(s,0,0,SW,Inches(0.22),BLUE); rect(s,0,Inches(7.28),SW,Inches(0.22),TEAL)
    org=_logo("tertiary-infotech-logo.png")
    if org: s.shapes.add_picture(org,Inches(0.85),Inches(0.7),height=Inches(1.05))
    # course badge (top-right)
    rect(s,Inches(11.0),Inches(0.72),Inches(1.55),Inches(1.0),BLUE)
    txt(s,Inches(11.0),Inches(0.80),Inches(1.55),Inches(0.5),[[("SY0-701",18,WHITE,True)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(11.0),Inches(1.26),Inches(1.55),Inches(0.4),[[("COMPTIA SECURITY+",7,WHITE,True)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(0.9),Inches(2.3),Inches(12),Inches(0.6),[[("COURSE SLIDES  ·  WSQ",16,BLUE,True)]])
    txt(s,Inches(0.9),Inches(2.85),Inches(12.0),Inches(1.9),[[(C.TITLE,40,INK,True)]])
    rect(s,Inches(0.92),Inches(4.35),Inches(2.4),Inches(0.06),TEAL)
    txt(s,Inches(0.9),Inches(4.65),Inches(12),Inches(1.4),
        [[(f"WSQ Course Code: {C.COURSE_CODE}",16,GREY,False)],
         [("Conducted by Tertiary Infotech Academy Pte Ltd  ·  UEN 201200696W",14,GREY,False)]],space=6)
    txt(s,Inches(0.9),Inches(6.5),Inches(12),Inches(0.4),[[(f"Version {C.VERSION}  ·  {C.VERSION_DATE}  ·  Trainer: {C.TRAINER}",12,GREY,False)]])
    txt(s,Inches(0.9),Inches(6.85),Inches(12),Inches(0.34),[[("© 2026 Tertiary Infotech Academy Pte Ltd. All rights reserved.  ·  www.tertiarycourses.com.sg",10,GREY,False)]])

def section(kicker,title,n,sub=""):
    s=slide(); rect(s,0,0,SW,SH,WHITE); rect(s,0,0,Inches(0.28),SH,BLUE)
    rect(s,Inches(0.85),Inches(2.5),Inches(0.14),Inches(2.0),TEAL)
    txt(s,Inches(1.25),Inches(2.55),Inches(11),Inches(0.6),[[(kicker,18,BLUE,True)]])
    txt(s,Inches(1.25),Inches(3.0),Inches(11.4),Inches(1.6),[[(title,40,INK,True)]])
    if sub: txt(s,Inches(1.27),Inches(4.55),Inches(11),Inches(0.8),[[(sub,16,GREY,False)]])
    txt(s,Inches(10.0),Inches(0.7),Inches(2.8),Inches(1.6),[[(n,72,RGBColor(0xE2,0xE8,0xF0),True)]],align=PP_ALIGN.RIGHT)
    footer(s)
def content(title,items,kicker=None,size=20):
    s=head(slide(),title,kicker); bullets(s,Inches(0.85),Inches(1.95),Inches(11.6),Inches(4.9),items,size=size); footer(s); return s
def two_col(title,left,right,kicker=None,lhead="",rhead=""):
    s=head(slide(),title,kicker)
    rect(s,Inches(0.85),Inches(1.95),Inches(5.7),Inches(4.7),LIGHT); rect(s,Inches(6.95),Inches(1.95),Inches(5.55),Inches(4.7),LIGHT)
    if lhead: txt(s,Inches(1.1),Inches(2.15),Inches(5.2),Inches(0.4),[[(lhead,16,BLUE,True)]])
    if rhead: txt(s,Inches(7.2),Inches(2.15),Inches(5.0),Inches(0.4),[[(rhead,16,TEAL,True)]])
    bullets(s,Inches(1.1),Inches(2.7),Inches(5.2),Inches(3.8),left,size=16)
    bullets(s,Inches(7.2),Inches(2.7),Inches(5.05),Inches(3.8),right,size=16,mcolor=TEAL); footer(s); return s
def cards3(title,cards,kicker):
    s=head(slide(),title,kicker); xs=[Inches(0.85),Inches(5.0),Inches(9.15)]
    for i,c in enumerate(cards[:3]):
        x=xs[i]; col=c[0]
        rect(s,x,Inches(1.95),Inches(3.65),Inches(4.7),LIGHT); rect(s,x,Inches(1.95),Inches(3.65),Inches(0.12),col)
        txt(s,x+Inches(0.25),Inches(2.2),Inches(3.2),Inches(0.6),[[(c[1],19,col,True)]])
        bullets(s,x+Inches(0.25),Inches(2.95),Inches(3.2),Inches(3.4),c[2],size=14,mcolor=col,gap=9)
    footer(s); return s
def big_statement(line1,line2,kicker,color=BLUE):
    s=slide(); rect(s,0,0,SW,SH,WHITE); rect(s,0,0,Inches(0.28),SH,color)
    txt(s,Inches(1.1),Inches(2.2),Inches(11),Inches(0.5),[[(kicker,16,color,True)]])
    txt(s,Inches(1.1),Inches(2.8),Inches(11.3),Inches(2.4),[[(line1,38,INK,True)]])
    if line2: txt(s,Inches(1.12),Inches(4.9),Inches(11),Inches(1.2),[[(line2,20,GREY,False)]])
    footer(s); return s
def tile_grid(title,items,kicker=None,cols=2,size=15,icons=None,accent=BLUE):
    s=head(slide(),title,kicker,kcolor=accent)
    n=len(items); rows=math.ceil(n/cols)
    X0=Inches(0.85); Y0=Inches(1.95); TOTW=Inches(11.63); AREAH=Inches(4.78)
    gx=Inches(0.3); gy=Inches(0.26)
    cw=int((TOTW-gx*(cols-1))/cols); ch=int((AREAH-gy*(rows-1))/rows)
    bd=Inches(0.6)
    for i,it in enumerate(items):
        r=i//cols; c=i%cols
        x=int(X0+(cw+gx)*c); y=int(Y0+(ch+gy)*r); col=PALETTE[i%len(PALETTE)]
        rect(s,x,y,cw,ch,LIGHT); rect(s,x,y,Inches(0.1),ch,col)
        oval(s,x+Inches(0.28),int(y+ch/2-bd/2),bd,bd,col)
        ic=icons[i] if icons else str(i+1)
        txt(s,x+Inches(0.28),int(y+ch/2-bd/2),bd,bd,[[(ic,19,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        tx=x+Inches(1.08); tw=cw-Inches(1.32)
        if isinstance(it,tuple):
            txt(s,tx,int(y+Inches(0.14)),tw,int(ch-Inches(0.2)),
                [[(it[0],size+2,INK,True)],[(it[1],size-2,GREY,False)]],anchor=MSO_ANCHOR.MIDDLE,space=3)
        else:
            txt(s,tx,int(y+Inches(0.1)),tw,int(ch-Inches(0.16)),[[(it,size,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s
def flow_h(title,steps,kicker=None,color=BLUE):
    s=head(slide(),title,kicker,kcolor=color)
    n=len(steps); X0=Inches(0.85); TOTW=Inches(11.63); gap=Inches(0.34)
    cw=int((TOTW-gap*(n-1))/n); y=Inches(2.55); ch=Inches(3.15); bd=Inches(0.82)
    for i,st in enumerate(steps):
        x=int(X0+(cw+gap)*i)
        rect(s,x,y,cw,ch,LIGHT); rect(s,x,y,cw,Inches(0.1),color)
        oval(s,int(x+cw/2-bd/2),int(y+Inches(0.42)),bd,bd,color)
        txt(s,int(x+cw/2-bd/2),int(y+Inches(0.42)),bd,bd,[[(str(i+1),30,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
        txt(s,x+Inches(0.16),int(y+Inches(1.55)),cw-Inches(0.32),int(ch-Inches(1.7)),[[(st,14,INK,False)]],align=PP_ALIGN.CENTER)
        if i<n-1:
            txt(s,int(x+cw-Inches(0.04)),int(y+ch/2-Inches(0.3)),int(gap+Inches(0.08)),Inches(0.6),
                [[("▶",15,color,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s
def trainer_slide(kicker,name,role,rows,initials,accent=BLUE):
    s=head(slide(),"About the Trainer",kicker,kcolor=accent)
    lx=Inches(0.85); lw=Inches(3.65)
    rect(s,lx,Inches(1.95),lw,Inches(4.7),LIGHT); rect(s,lx,Inches(1.95),lw,Inches(0.12),accent)
    bd=Inches(1.7); ax=int(lx+(lw-bd)/2)
    oval(s,ax,Inches(2.5),bd,bd,accent)
    txt(s,ax,Inches(2.5),bd,bd,[[(initials,44,WHITE,True)]],align=PP_ALIGN.CENTER,anchor=MSO_ANCHOR.MIDDLE)
    txt(s,lx+Inches(0.15),Inches(4.55),lw-Inches(0.3),Inches(0.6),[[(name,21,INK,True)]],align=PP_ALIGN.CENTER)
    txt(s,lx+Inches(0.15),Inches(5.2),lw-Inches(0.3),Inches(1.2),[[(role,13,GREY,False)]],align=PP_ALIGN.CENTER)
    rx=Inches(4.9); rw=Inches(7.6); ry=Inches(1.95); rh=Inches(4.7)
    n=len(rows); gy=Inches(0.2); th=int((rh-gy*(n-1))/n)
    for i,(label,val) in enumerate(rows):
        y=int(ry+(th+gy)*i); col=PALETTE[i%len(PALETTE)]
        rect(s,rx,y,rw,th,LIGHT); rect(s,rx,y,Inches(0.1),th,col)
        vruns=[(val,14,INK,False)] if val else [("____________________________________________",13,LINE,False)]
        txt(s,rx+Inches(0.32),y,rw-Inches(0.6),th,
            [[(label.upper(),11,col,True)],vruns],anchor=MSO_ANCHOR.MIDDLE,space=3)
    footer(s); return s
def access_labs():
    """'Access the Hands-On Labs' slide with a clickable GitHub repo link."""
    s=head(slide(),"Access the Hands-On Labs","36 LABS · FREE IN YOUR BROWSER",kcolor=VIOLET)
    txt(s,Inches(0.85),Inches(1.95),Inches(11.6),Inches(0.5),
        [[("All 36 labs live in the course GitHub repository: ",16,INK,False)]])
    # clickable repo URL
    tb=s.shapes.add_textbox(Inches(0.85),Inches(2.4),Inches(11.6),Inches(0.5)); tf=tb.text_frame; tf.word_wrap=True
    p=tf.paragraphs[0]; r=p.add_run(); r.text=C.REPO_URL; r.font.size=Pt(18); r.font.bold=True
    r.font.name="Arial"; r.font.color.rgb=BLUE; r.hyperlink.address=C.REPO_URL
    # two option cards
    xs=[Inches(0.85),Inches(6.95)]
    cards=[(BLUE,"Option A — Clone",["Open a terminal / Killercoda","git clone the repository","cd into labs/ and follow along"],
            f"git clone {C.REPO_URL}.git"),
           (TEAL,"Option B — Download ZIP",["On GitHub click the green Code button","Choose Download ZIP","Unzip and open the labs/ folder"],"")]
    for i,(col,ttl,items,code) in enumerate(cards):
        x=xs[i]; rect(s,x,Inches(3.15),Inches(5.55),Inches(2.9),LIGHT); rect(s,x,Inches(3.15),Inches(5.55),Inches(0.12),col)
        txt(s,x+Inches(0.3),Inches(3.35),Inches(5),Inches(0.5),[[(ttl,18,col,True)]])
        bullets(s,x+Inches(0.3),Inches(3.95),Inches(5),Inches(1.4),items,size=14,mcolor=col,gap=7)
        if code:
            rect(s,x+Inches(0.3),Inches(5.45),Inches(4.95),Inches(0.5),RGBColor(0x0B,0x12,0x20))
            txt(s,x+Inches(0.45),Inches(5.5),Inches(4.7),Inches(0.4),[[("$ "+code,11,RGBColor(0x9C,0xDC,0xFE),False)]],anchor=MSO_ANCHOR.MIDDLE)
    txt(s,Inches(0.85),Inches(6.25),Inches(11.6),Inches(0.5),
        [[("Every lab runs free in the browser on the Killercoda Ubuntu Playground — no local install required.",13,GREY,False)]])
    footer(s); return s
def activity_overview(tag,title,desc,build,services,kicker):
    s=head(slide(),title,kicker,kcolor=TEAL)
    rect(s,Inches(0.85),Inches(1.85),Inches(1.7),Inches(0.5),TEAL)
    txt(s,Inches(0.85),Inches(1.9),Inches(1.7),Inches(0.4),[[(tag,16,WHITE,True)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(0.85),Inches(2.55),Inches(11.7),Inches(1.6),[[(desc,21,INK,False)]])
    rect(s,Inches(0.85),Inches(4.3),Inches(11.7),Inches(2.0),LIGHT)
    txt(s,Inches(1.1),Inches(4.5),Inches(11),Inches(0.4),[[("You'll build",14,BLUE,True)]])
    txt(s,Inches(1.1),Inches(4.9),Inches(11),Inches(0.6),[[(build,18,INK,True)]])
    txt(s,Inches(1.1),Inches(5.6),Inches(11.2),Inches(0.6),[[("Tools:  ",13,GREY,True),(services,13,GREY,False)]]); footer(s); return s
def step_slide(kicker,act_title,n,total,text,cmd=""):
    s=head(slide(),act_title,kicker,TEAL)
    oval(s,Inches(0.85),Inches(2.5),Inches(1.4),Inches(1.4),TEAL)
    txt(s,Inches(0.85),Inches(2.74),Inches(1.4),Inches(0.9),[[(str(n),38,WHITE,True)]],align=PP_ALIGN.CENTER)
    txt(s,Inches(0.95),Inches(1.95),Inches(11),Inches(0.4),[[(f"STEP {n} OF {total}",13,GREY,True)]])
    txt(s,Inches(2.55),Inches(2.4),Inches(10.1),Inches(1.3),[[(text,23,INK,False)]],anchor=MSO_ANCHOR.MIDDLE)
    if cmd:
        rect(s,Inches(2.55),Inches(4.15),Inches(10.1),Inches(0.95),RGBColor(0x0B,0x12,0x20))
        txt(s,Inches(2.8),Inches(4.28),Inches(9.7),Inches(0.7),[[("$ "+cmd,13,RGBColor(0x9C,0xDC,0xFE),False)]],anchor=MSO_ANCHOR.MIDDLE)
    footer(s); return s
def test_slide(act_title,text,kicker):
    s=head(slide(),act_title,kicker,TEAL)
    rect(s,Inches(0.85),Inches(2.3),Inches(11.7),Inches(2.6),RGBColor(0xE8,0xF7,0xEE))
    txt(s,Inches(1.2),Inches(2.6),Inches(11),Inches(0.5),[[("✅  Test it",20,RGBColor(0x12,0x7A,0x3E),True)]])
    txt(s,Inches(1.2),Inches(3.3),Inches(11),Inches(1.4),[[(text,18,INK,False)]]); footer(s); return s
def brk(kind,dur,color=AMBER):
    s=slide(); rect(s,0,0,SW,SH,WHITE)
    rect(s,0,0,SW,Inches(0.22),color); rect(s,0,Inches(7.28),SW,Inches(0.22),color)
    rect(s,Inches(5.4),Inches(2.35),Inches(2.53),Inches(0.1),color)
    txt(s,0,Inches(2.75),SW,Inches(1.2),[[(kind,48,INK,True)]],align=PP_ALIGN.CENTER)
    txt(s,0,Inches(4.05),SW,Inches(0.8),[[(dur,22,color,True)]],align=PP_ALIGN.CENTER); PAGE["n"]+=1

# ============================================================ BUILD
cover()

# ---------------- ADMIN ----------------
section("COURSE ADMINISTRATION","Welcome & Housekeeping","")
content("Digital Attendance (Mandatory)",[
 "It is mandatory to take the AM, PM and Assessment digital attendance for WSQ-funded courses.",
 "The trainer/administrator displays the digital attendance QR code from the SSG portal.",
 "Scan the QR code with your mobile phone camera and submit your attendance.",
 "A minimum of 75% attendance is required to be eligible for assessment and funding."],kicker="TRAQOM · SSG DIGITAL ATTENDANCE")
trainer_slide("YOUR TRAINER · GENERAL","Your Trainer","General Trainer template —\nto be completed by the trainer",
 [("Name",""),("Title / Designation",""),("Qualifications",""),
  ("Areas of expertise",""),("Training & industry experience",""),("Contact","")],
 initials="?",accent=GREY)
trainer_slide("YOUR TRAINER",C.TRAINER,"Principal Trainer\nTertiary Infotech Academy Pte. Ltd.",
 [("Role","Principal Trainer, Tertiary Infotech Academy Pte. Ltd."),
  ("Certification","CompTIA Security+ / cybersecurity — defensive & offensive security."),
  ("Delivers","WSQ courses on cybersecurity, cloud, networking & software engineering."),
  ("Founder","Founder and lead instructor at Tertiary Infotech / Tertiary Courses.")],
 initials="AA",accent=BLUE)
content("Let's Know Each Other",[
 "Your name and organisation / role.",
 "Your experience with IT security, networking or Linux (if any).",
 "What you want to be able to do — or which exam domain you want to master — after this course."],kicker="ICE-BREAKER")
tile_grid("Ground Rules",[
 "Set your mobile phone to silent mode.","Participate actively — no question is too small.",
 "Mutual respect: agree to disagree.","One conversation at a time.",
 "Be punctual; return from breaks on time.","75% attendance is required."],
 kicker="HOUSEKEEPING",cols=2,size=15)
content("LMS / TMS",[
 "Access your course materials, attendance and assessment on the LMS/TMS portal.",
 "Portal: https://lms-tms.tertiaryinfotech.com",
 "Download the slides and Learner Guide for reference during the open-book assessment."],kicker="COURSE PORTAL")
# Lesson plan overview (3 days)
two_col("Lesson Plan — 3 Days, 8 hours/day",[
 (f"Day 1 — {C.DAY_THEMES[1]}",0),
 ("Domain 1: General Security Concepts (Labs 1–5)",1),
 ("Domain 2: Threats & Vulnerabilities (Labs 6–13)",1),
 (f"Day 2 — {C.DAY_THEMES[2]}",0),
 ("Domain 3: Security Architecture (Labs 14–20)",1),
 ("Domain 4: Security Operations begins (Labs 21–26)",1)],
 [(f"Day 3 — {C.DAY_THEMES[3]}",0),
 ("Domain 4: Security Operations (Labs 27–30)",1),
 ("Domain 5: Program Management (Labs 31–36)",1),
 ("Final Assessment (WA + PP)",1),
 ("Daily timing",0),
 ("9:30am–6:30pm · 1-hour lunch · tea breaks within",1)],
 kicker="SCHEDULE",lhead="Days 1–2",rhead="Day 3 & timing")
tile_grid("Learning Outcomes",[
 ("LO1 · General Concepts","Controls, CIA & AAA, hashing, encryption and PKI."),
 ("LO2 · Threats & Vulns","Phishing, password & web attacks, malware, scanning, hardening."),
 ("LO3 · Architecture","Segmentation, firewalls, WAF, VPN, encryption, TLS, backup."),
 ("LO4 · Operations","Hardening, MFA, SIEM, IDS, EDR, FIM, vuln mgmt, IR, SOAR."),
 ("LO5 · Program Mgmt","Risk & CVSS, recon, pentesting, compliance, awareness.")],
 kicker="WHAT YOU'LL ACHIEVE",cols=2,size=14)
content("Briefing for Assessment",[
 "Place phones and other materials under the table or on the floor.",
 "No photos or recording of assessment scripts.","No discussion during the assessment.",
 "Use a black/blue pen for hard-copy assessments.","No liquid paper / correction tape.",
 "Scripts are collected when time is up."])
content("Assessment",[
 C.ASSESSMENT["written"], C.ASSESSMENT["practical"],
 "Format: Open Book — slides, Learner Guide and approved materials only.",
 C.ASSESSMENT["note"],"An appeal process is available if required."],kicker="FINAL ASSESSMENT")
flow_h("Assessment Flow",[
 "Take the Assessment digital attendance (TRAQOM · SSG QR on the LMS)",
 "Sit WA (SAQ) — 6 questions, open book",
 "Sit PP — 5 hands-on tasks, open book",
 "Submit your answers on the LMS",
 "Sign the Assessment Summary Record"],kicker="ON ASSESSMENT DAY")

# ---------------- CORE CONCEPTS ----------------
section("CORE CONCEPTS","Security Fundamentals for the SY0-701","")
tile_grid("What is CompTIA Security+?",[
 ("Vendor-neutral baseline","The industry standard entry certification for cybersecurity roles."),
 ("SY0-701 domains","Five domains from general concepts to program management."),
 ("Hands-on and policy","Tests both technical skills and risk / governance knowledge."),
 ("The security role","Assess posture, apply controls, monitor, respond and manage risk.")],
 kicker="OVERVIEW",cols=2,size=15)
cards3("The Five Exam Domains (SY0-701)",[
 (BLUE,"Concepts & Threats",["1. General Security Concepts — 12%","2. Threats, Vulns & Mitigations — 22%"]),
 (TEAL,"Build & Operate",["3. Security Architecture — 18%","4. Security Operations — 28%"]),
 (VIOLET,"Govern",["5. Security Program Management & Oversight — 20%","Passing score: 750 / 900"])],kicker="BLUEPRINT")
big_statement("Defence in depth: layer controls so no single failure is fatal.","Confidentiality, integrity and availability are the goals — every control in this course serves one of them.","WHY IT MATTERS",color=BLUE)
tile_grid("The CIA Triad & AAA",[
 ("Confidentiality","Only authorised parties can read the data — encryption, access control."),
 ("Integrity","Data is not altered undetected — hashing, digital signatures."),
 ("Availability","Systems and data are there when needed — redundancy, backup, DoS defence."),
 ("AAA","Authentication proves who you are, authorization what you may do, accounting logs it.")],
 kicker="FOUNDATIONS",cols=2,size=15)
content("Your Lab Workbench — Killercoda",[
 "Every hands-on lab runs in a free, disposable Ubuntu VM in your browser.",
 f"Open the playground: {C.KILLERCODA_URL}",
 "You have a root shell with internet access — install tools with apt or a single binary.",
 "Reset the playground between labs that change firewall, PAM or routing state."],kicker="YOUR WORKBENCH")
access_labs()

# ---------------- TOPICS + ACTIVITIES ----------------
TOPIC_ACTS = {t["num"]: [a for a in ACTIVITIES if a["topic"]==t["num"]] for t in C.TOPICS}
CARD_COLORS=[BLUE,TEAL,VIOLET]
for t in C.TOPICS:
    section(f"DOMAIN {t['code']}", t["title"], t["code"], t["subtitle"])
    tile_grid(f"Key Concepts — {t['title']}", t["concepts"],
              kicker=f"EXAM WEIGHTING {t['weighting']}", cols=2, size=14)
    acts=TOPIC_ACTS[t["num"]]
    # a card summary of the labs in this domain
    third=(len(acts)+2)//3
    groups=[acts[i:i+third] for i in range(0,len(acts),third)][:3]
    while len(groups)<3: groups.append([])
    cards=[]
    for gi,g in enumerate(groups):
        cards.append((CARD_COLORS[gi], f"Labs {g[0]['num']}–{g[-1]['num']}" if g else "—",
                      [a["title"] for a in g] if g else ["—"]))
    cards3(f"Hands-On Labs — {t['title']}", cards, kicker="WHAT YOU'LL DO")
    # per activity
    for a in acts:
        activity_overview(f"LAB {a['num']}", a["title"], a["desc"], a["build"], a["services"], kicker=f"DOMAIN {t['code']} · HANDS-ON")
        steps=a["steps"]; total=len(steps)
        for i,(instr,cmd) in enumerate(steps,1):
            step_slide(f"LAB {a['num']} · {a['title'][:34]}", a["title"], i, total, instr, cmd)
        test_slide(a["title"], a["test"], kicker=f"LAB {a['num']} · VERIFY")
    # domain recap
    content(f"Recap — {t['title']}",
            ["You can now: "+a["objective"] for a in {x["objective"]:x for x in acts}.values()][:8],
            kicker="DOMAIN RECAP", size=16)

# ---------------- CLOSE ----------------
section("WRAP-UP","Course Summary & Next Steps","")
tile_grid("What You Achieved",[
 ("General Concepts","Categorised controls; applied CIA, AAA, hashing, encryption and PKI."),
 ("Threats & Vulns","Analysed phishing, cracked passwords, exploited SQLi/XSS, scanned and hardened."),
 ("Architecture","Segmented networks; deployed firewalls, WAF, VPN, disk & TLS encryption and backup."),
 ("Operations","Hardened access; ran SIEM, IDS, EDR, FIM, vuln management, IR and automation."),
 ("Program Mgmt","Scored risk with CVSS; ran recon, pentesting, compliance and awareness testing.")],
 kicker="LEARNING OUTCOMES",cols=2,size=14)
content("Preparing for the SY0-701 Exam",[
 "Redo every lab until the commands and concepts are automatic.",
 "Review the 'What you learned' notes in each lab and the Learner Guide.",
 "Memorise the control categories/types, the CIA triad and the CVSS metrics.",
 "The SY0-701 exam is up to 90 questions in 90 minutes; passing score is 750/900.",
 "Book the exam through Pearson VUE from your CompTIA account."],kicker="NEXT STEPS")
content("Practice Exam",[
 "Sharpen your exam readiness with the Tertiary Infotech Security+ practice exam.",
 f"Practice exam: {C.PRACTICE_EXAM_URL}",
 "Attempt it under timed conditions and review every explanation.",
 "Revisit any lab whose topic you miss, then re-take the practice exam."],kicker="TEST YOURSELF")
content("Courseware & Assessment on the LMS",[
 "Download the slides and Learner Guide from the LMS at https://lms-tms.tertiaryinfotech.com/.",
 "Complete the Written Assessment (WA) and Practical Performance (PP) on the LMS.",
 "Keep the courseware open during the assessment — it is open book.",
 "Remember to take the Assessment digital attendance (TRAQOM)."],kicker="LMS / TMS")
# ---- assessment admin repeated at END (house rule) ----
content("Assessment",[
 "Written Assessment (SAQ) — 6 questions, 1 hour.  Practical Performance (PP) — 5 tasks, 1 hour.",
 "Open book: slides, Learner Guide and approved materials only.",
 "Remember to take the Assessment digital attendance (TRAQOM).",
 "Submit your completed answers on the LMS at https://lms-tms.tertiaryinfotech.com/."],kicker="WRAP-UP")
flow_h("Assessment Flow",[
 "Take the Assessment digital attendance (TRAQOM · SSG QR on the LMS)",
 "Sit WA (SAQ) — 6 questions, open book",
 "Sit PP — 5 hands-on tasks, open book",
 "Submit your answers on the LMS",
 "Sign the Assessment Summary Record"],kicker="ON ASSESSMENT DAY")
content("Digital Attendance (Mandatory)",[
 "It is mandatory to take the AM, PM and Assessment digital attendance for WSQ-funded courses.",
 "The trainer/administrator displays the digital attendance QR code from the SSG portal.",
 "Scan the QR code with your mobile phone camera and submit your attendance.",
 "A minimum of 75% attendance is required to be eligible for assessment and funding."],kicker="TRAQOM · SSG DIGITAL ATTENDANCE")
big_statement("Thank You!","You are now ready to secure Linux systems — and to sit the CompTIA Security+ SY0-701 exam.","STAY SECURE",color=TEAL)

OUT=os.path.join(REPO,f"{C.SHORT_TITLE}-{C.VERSION}.pptx")
prs.save(OUT)
print(f"Saved {OUT}  ({len(prs.slides._sldIdLst)} slides)")
